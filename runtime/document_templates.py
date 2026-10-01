"""Safe DOCX template metadata, placeholder validation, and rendering."""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import uuid
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


PLACEHOLDER_RE = re.compile(r"{{\s*([#/]?[^{}]+?)\s*}}")
SUPPORTED_XML_PARTS = ("word/document.xml", "word/header", "word/footer")
IMAGE_SLOT_NAMES = frozenset({"image_surface", "image_mask", "image_measurement"})
DEFAULT_IMAGE_WIDTH_INCHES = 4.8
DEFAULT_IMAGE_HEIGHT_INCHES = 2.6


@dataclass
class TemplateMetadata:
    template_id: str
    path: str
    name: str
    version: int
    sha256: str
    placeholders: list[str]
    status: str
    error: str = ""


class TemplateValidationError(ValueError):
    pass


@dataclass
class ImageSlotStatus:
    slot: str
    source_path: str = ""
    status: str = "skipped"
    error: str = ""


@dataclass
class DocumentRenderResult:
    path: Path
    image_slots: list[ImageSlotStatus]
    unreplaced_placeholders: list[str]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _xml_parts(archive: zipfile.ZipFile) -> list[str]:
    return [name for name in archive.namelist() if name == "word/document.xml" or name.startswith("word/header") or name.startswith("word/footer")]


def scan_template(path: str | Path, *, allowed_fields: set[str] | None = None) -> TemplateMetadata:
    source = Path(path).resolve()
    if source.suffix.lower() != ".docx":
        raise TemplateValidationError("模板必须是 DOCX 文件")
    digest = _sha256(source)
    template_id = hashlib.sha256(f"{source}:{digest}".encode()).hexdigest()[:20]
    placeholders: list[str] = []
    try:
        with zipfile.ZipFile(source) as archive:
            for part in _xml_parts(archive):
                text = archive.read(part).decode("utf-8")
                placeholders.extend(match.group(1).strip() for match in PLACEHOLDER_RE.finditer(text))
                try:
                    root = ElementTree.fromstring(text)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    for paragraph in root.findall(".//w:p", ns):
                        combined = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
                        placeholders.extend(match.group(1).strip() for match in PLACEHOLDER_RE.finditer(combined))
                except ElementTree.ParseError:
                    pass
    except (OSError, zipfile.BadZipFile, UnicodeDecodeError) as exc:
        return TemplateMetadata(template_id, str(source), source.name, 1, digest, [], "invalid", f"模板读取失败：{type(exc).__name__}")
    errors: list[str] = []
    stack = 0
    for placeholder in placeholders:
        if placeholder.startswith("#each "):
            stack += 1
        elif placeholder == "/each":
            stack -= 1
            if stack < 0:
                errors.append("发现未匹配的 {{/each}}")
                stack = 0
        elif placeholder.startswith("/"):
            errors.append(f"不支持的结束占位符：{placeholder}")
        elif allowed_fields and placeholder not in allowed_fields:
            errors.append(f"未知占位符：{placeholder}")
    if stack:
        errors.append("发现未闭合的 findings 列表占位符")
    return TemplateMetadata(template_id, str(source), source.name, 1, digest, sorted(set(placeholders)), "invalid" if errors else "ready", "；".join(errors))


def register_managed_template(path: str | Path, managed_dir: str | Path) -> TemplateMetadata:
    """Validate a template and register a project-owned copy, never an external path."""
    source = Path(path).resolve()
    metadata = scan_template(source)
    if metadata.status != "ready":
        return metadata
    destination_dir = Path(managed_dir).resolve()
    destination_dir.mkdir(parents=True, exist_ok=True)
    candidate = destination_dir / source.name
    if candidate.exists() and _sha256(candidate) != metadata.sha256:
        candidate = destination_dir / f"{source.stem}_{metadata.sha256[:12]}{source.suffix.lower()}"
    if candidate.resolve() != source:
        shutil.copy2(source, candidate)
    return scan_template(candidate)


def _replace_text(xml: str, values: dict[str, Any]) -> str:
    def replace(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        if key.startswith("#each ") or key == "/each":
            return match.group(0)
        return str(values.get(key, match.group(0)))

    return PLACEHOLDER_RE.sub(replace, xml)


def _replace_cross_run_placeholders(xml: str, values: dict[str, Any]) -> str:
    """Replace placeholders split across Word runs while keeping the paragraph XML valid."""
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        root = ElementTree.fromstring(xml)
    except ElementTree.ParseError:
        return xml
    changed = False
    for paragraph in root.findall(".//w:p", ns):
        text_nodes = paragraph.findall(".//w:t", ns)
        if len(text_nodes) < 2:
            continue
        combined = "".join(node.text or "" for node in text_nodes)
        if "{{" not in combined:
            continue
        replaced = _replace_text(combined, values)
        text_nodes[0].text = replaced
        for node in text_nodes[1:]:
            node.text = ""
        changed = True
    if not changed:
        return xml
    return ElementTree.tostring(root, encoding="unicode")


def _iter_paragraphs(container: Any):
    """Yield paragraphs in a document section or table hierarchy exactly once."""
    seen: set[int] = set()

    def visit(current: Any):
        for paragraph in getattr(current, "paragraphs", []):
            element_id = id(paragraph._p)
            if element_id not in seen:
                seen.add(element_id)
                yield paragraph
        for table in getattr(current, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from visit(cell)

    yield from visit(container)


def _image_slot_paragraphs(document: Any):
    seen: set[int] = set()
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in _iter_paragraphs(container):
            element_id = id(paragraph._p)
            if element_id not in seen:
                seen.add(element_id)
                yield paragraph


def _replace_image_slots(
    document_path: Path,
    image_slots: dict[str, str | Path | None],
    *,
    width_inches: float,
    height_inches: float,
) -> list[ImageSlotStatus]:
    """Replace standalone image placeholders with reliable inline DOCX pictures."""
    try:
        from docx import Document
        from docx.image.image import Image
        from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover - depends on deployment extras
        return [
            ImageSlotStatus(slot=slot, status="failed", error="缺少 python-docx 图像嵌入组件")
            for slot in image_slots
        ]

    document = Document(document_path)
    located: dict[str, list[Any]] = {slot: [] for slot in image_slots}
    for paragraph in _image_slot_paragraphs(document):
        text = "".join(run.text for run in paragraph.runs).strip()
        matched = PLACEHOLDER_RE.fullmatch(text)
        if matched and matched.group(1).strip() in located:
            located[matched.group(1).strip()].append(paragraph)

    statuses: list[ImageSlotStatus] = []
    for slot, raw_source in image_slots.items():
        paragraphs = located.get(slot, [])
        source = Path(raw_source) if raw_source else None
        source_text = str(source) if source is not None else ""
        if not paragraphs:
            statuses.append(ImageSlotStatus(slot, source_text, "skipped", "模板中未找到图片槽位"))
            continue
        if source is None or not source.is_file():
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.text = ""
                paragraph.add_run("（未提供图像）")
            statuses.append(ImageSlotStatus(slot, source_text, "skipped", "图像文件不可用"))
            continue
        try:
            image = Image.from_file(str(source))
            width_limit = int(Inches(width_inches))
            height_limit = int(Inches(height_inches))
            scale = min(width_limit / image.width, height_limit / image.height)
            image_width = int(image.width * scale)
            image_height = int(image.height * scale)
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.text = ""
                paragraph.add_run().add_picture(str(source), width=image_width, height=image_height)
                paragraph.paragraph_format.keep_with_next = True
            statuses.append(ImageSlotStatus(slot, source_text, "inserted"))
        except Exception as exc:
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.text = ""
                paragraph.add_run("（图像嵌入失败）")
            statuses.append(ImageSlotStatus(slot, source_text, "failed", f"{type(exc).__name__}: {exc}"))
    document.save(document_path)
    return statuses


def _apply_output_text_cleanup(document_path: Path, replacements: dict[str, str]) -> None:
    """Apply approved output-only wording fixes without changing the source template."""
    if not replacements:
        return
    cleanup_path = document_path.with_name(
        f".{document_path.stem}.{uuid.uuid4().hex}.cleanup{document_path.suffix}"
    )
    try:
        with zipfile.ZipFile(document_path) as archive, zipfile.ZipFile(
            cleanup_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as output:
            xml_parts = set(_xml_parts(archive))
            for info in archive.infolist():
                data = archive.read(info.filename)
                if info.filename in xml_parts:
                    root = ElementTree.fromstring(data.decode("utf-8"))
                    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    for paragraph in root.findall(".//w:p", namespace):
                        text_nodes = paragraph.findall(".//w:t", namespace)
                        if not text_nodes:
                            continue
                        original = "".join(node.text or "" for node in text_nodes)
                        cleaned = original
                        for source, target in replacements.items():
                            cleaned = cleaned.replace(source, target)
                        if cleaned == original:
                            continue
                        text_nodes[0].text = cleaned
                        for node in text_nodes[1:]:
                            node.text = ""
                    data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(info, data)
        os.replace(cleanup_path, document_path)
    except (OSError, UnicodeDecodeError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        try:
            cleanup_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise TemplateValidationError(f"DOCX 文本清理失败：{type(exc).__name__}") from exc


def render_docx_with_status(
    template: str | Path,
    output: str | Path,
    values: dict[str, Any],
    *,
    image_slots: dict[str, str | Path | None] | None = None,
    optional_image_slots: set[str] | None = None,
    output_text_cleanup: dict[str, str] | None = None,
    image_width_inches: float = DEFAULT_IMAGE_WIDTH_INCHES,
    image_height_inches: float = DEFAULT_IMAGE_HEIGHT_INCHES,
    required: bool = False,
) -> DocumentRenderResult:
    source = Path(template)
    target = Path(output)
    resolved_image_slots = dict(image_slots or {})
    optional_slots = set(optional_image_slots or set())
    if not set(resolved_image_slots).issubset(IMAGE_SLOT_NAMES):
        unsupported = sorted(set(resolved_image_slots) - IMAGE_SLOT_NAMES)
        raise TemplateValidationError(f"不支持的图片槽位：{', '.join(unsupported)}")
    if not optional_slots.issubset(resolved_image_slots):
        unsupported = sorted(optional_slots - set(resolved_image_slots))
        raise TemplateValidationError(f"可选图片槽位未包含在渲染槽位中：{', '.join(unsupported)}")
    nested_fields = {key for item in values.get("findings", []) if isinstance(item, dict) for key in item}
    metadata = scan_template(
        source,
        allowed_fields=set(values) | nested_fields | set(resolved_image_slots) | {"findings", "/each"},
    )
    if metadata.status != "ready":
        raise TemplateValidationError(metadata.error)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.stem}.{uuid.uuid4().hex}.tmp{target.suffix}")
    findings = values.get("findings")
    try:
        with zipfile.ZipFile(source) as archive, zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as out:
            for info in archive.infolist():
                data = archive.read(info.filename)
                if info.filename in _xml_parts(archive):
                    xml = data.decode("utf-8")
                    if isinstance(findings, list):
                        block = re.compile(r"{{\s*#each findings\s*}}(.*?){{\s*/each\s*}}", re.DOTALL)

                        def render_block(match: re.Match[str]) -> str:
                            return "".join(_replace_text(match.group(1), item if isinstance(item, dict) else {}) for item in findings)

                        xml = block.sub(render_block, xml)
                    xml = _replace_cross_run_placeholders(xml, values)
                    xml = _replace_text(xml, values)
                    data = xml.encode("utf-8")
                out.writestr(info, data)
        _apply_output_text_cleanup(temporary, dict(output_text_cleanup or {}))
        statuses = _replace_image_slots(
            temporary,
            resolved_image_slots,
            width_inches=max(1.0, float(image_width_inches)),
            height_inches=max(1.0, float(image_height_inches)),
        ) if resolved_image_slots else []
        failed_slots = [item.slot for item in statuses if item.status == "failed"]
        required_skips = [
            item.slot
            for item in statuses
            if required and item.slot not in optional_slots and item.status != "inserted"
        ]
        if failed_slots or required_skips:
            slots = failed_slots or required_skips
            raise TemplateValidationError(f"图片槽位未完成：{', '.join(slots)}")
        unreplaced = scan_template(temporary).placeholders
        if unreplaced:
            raise TemplateValidationError(f"DOCX 中仍有未替换占位符：{', '.join(unreplaced)}")
        os.replace(temporary, target)
        return DocumentRenderResult(target, statuses, unreplaced)
    except Exception:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def render_docx(
    template: str | Path,
    output: str | Path,
    values: dict[str, Any],
    *,
    image_slots: dict[str, str | Path | None] | None = None,
    output_text_cleanup: dict[str, str] | None = None,
    image_width_inches: float = DEFAULT_IMAGE_WIDTH_INCHES,
    image_height_inches: float = DEFAULT_IMAGE_HEIGHT_INCHES,
    required: bool = False,
) -> Path:
    """Render a template and return its output path for backwards compatibility."""
    return render_docx_with_status(
        template,
        output,
        values,
        image_slots=image_slots,
        output_text_cleanup=output_text_cleanup,
        image_width_inches=image_width_inches,
        image_height_inches=image_height_inches,
        required=required,
    ).path


def metadata_dict(metadata: TemplateMetadata) -> dict[str, Any]:
    return asdict(metadata)
