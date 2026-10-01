from __future__ import annotations

import zipfile
from pathlib import Path


_PNG_BYTES = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000c49444154789c63f8cfc0000003010100c9fe92ef0000000049454e44ae426082"
)


def test_render_docx_embeds_inspection_images_and_clears_slots(tmp_path: Path) -> None:
    from docx import Document

    from runtime.document_templates import render_docx_with_status

    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{title}}")
    document.add_paragraph("{{image_surface}}")
    document.add_paragraph("{{image_mask}}")
    document.add_paragraph("{{image_measurement}}")
    document.save(template)
    image = tmp_path / "evidence.png"
    image.write_bytes(_PNG_BYTES)

    output = tmp_path / "report.docx"
    result = render_docx_with_status(
        template,
        output,
        {"title": "损伤分析报告"},
        image_slots={
            "image_surface": image,
            "image_mask": image,
            "image_measurement": image,
        },
    )

    assert result.path == output
    assert {item.status for item in result.image_slots} == {"inserted"}
    assert result.unreplaced_placeholders == []
    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        assert "{{image_" not in document_xml
        assert len([name for name in archive.namelist() if name.startswith("word/media/")]) == 1


def test_render_docx_marks_missing_image_as_skipped_without_placeholder(tmp_path: Path) -> None:
    from docx import Document

    from runtime.document_templates import render_docx_with_status

    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{image_surface}}")
    document.save(template)

    output = tmp_path / "report.docx"
    result = render_docx_with_status(
        template,
        output,
        {},
        image_slots={"image_surface": tmp_path / "missing.png"},
    )

    assert result.image_slots[0].status == "skipped"
    assert result.unreplaced_placeholders == []
    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "{{image_surface}}" not in document_xml
    assert "未提供图像" in document_xml


def test_render_docx_applies_output_only_text_cleanup(tmp_path: Path) -> None:
    from docx import Document

    from runtime.document_templates import render_docx_with_status

    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("（示例）宽度 {{width}}mm；面积 {{ratio}}%%")
    document.save(template)

    output = tmp_path / "report.docx"
    render_docx_with_status(
        template,
        output,
        {"width": "12 px（无物理标定）", "ratio": "12.00"},
        output_text_cleanup={
            "（示例）": "",
            "（无物理标定）mm": "（无物理标定）",
            "%%": "%",
        },
    )

    assert Document(output).paragraphs[0].text == "宽度 12 px（无物理标定）；面积 12.00%"


def test_render_docx_text_cleanup_does_not_import_docx_and_handles_split_runs(
    tmp_path: Path, monkeypatch
) -> None:
    import builtins

    from docx import Document

    from runtime.document_templates import render_docx_with_status

    template = tmp_path / "template.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("宽度 12 px（无物理")
    paragraph.add_run("标定）mm；面积 12.00%%")
    document.save(template)

    original_import = builtins.__import__

    def reject_docx_import(name, *args, **kwargs):
        if name == "docx" or name.startswith("docx."):
            raise ImportError("python-docx unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", reject_docx_import)
    output = tmp_path / "report.docx"
    render_docx_with_status(
        template,
        output,
        {},
        output_text_cleanup={"（无物理标定）mm": "（无物理标定）", "%%": "%"},
    )

    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "无物理标定）mm" not in document_xml
    assert "12.00%%" not in document_xml
    assert "无物理标定）" in document_xml
    assert "12.00%" in document_xml


def test_render_docx_failure_preserves_existing_report_and_removes_temporary_file(tmp_path: Path) -> None:
    from docx import Document
    import pytest

    from runtime.document_templates import TemplateValidationError, render_docx_with_status

    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{image_surface}}")
    document.save(template)
    corrupt_image = tmp_path / "corrupt.png"
    corrupt_image.write_bytes(b"not an image")
    output = tmp_path / "report.docx"
    existing_contents = b"previous-good-report"
    output.write_bytes(existing_contents)

    with pytest.raises(TemplateValidationError, match="图片槽位未完成"):
        render_docx_with_status(template, output, {}, image_slots={"image_surface": corrupt_image})

    assert output.read_bytes() == existing_contents
    assert not list(tmp_path.glob(".report.*.tmp.docx"))


def test_report_image_slots_pick_first_complete_successful_result(tmp_path: Path) -> None:
    import runtime.damage_workflow_gui as gui

    paths = [tmp_path / name for name in ("source.jpg", "mask.jpg", "legacy-measurement.png")]
    for path in paths:
        path.write_bytes(b"image")
    slots = gui._report_image_slots_from_summary(
        {
            "results": [
                {"status": "success", "image_path": str(paths[0])},
                {
                    "status": "success",
                    "image_path": str(paths[0]),
                    "overlay_path": str(paths[1]),
                    "measurement_path": str(paths[2]),
                },
            ]
        }
    )

    assert slots == {
        "image_surface": str(paths[0]),
        "image_mask": str(paths[1]),
        "image_measurement": None,
    }


def test_detection_finding_is_simplified_for_report_summary() -> None:
    import runtime.damage_workflow_gui as gui

    simplified = gui._simplify_detection_finding({
        "index": 2,
        "class_id": 6,
        "class_name": "Structural crack",
        "detection_confidence": 0.91,
        "area": {"area_ratio": 0.2},
        "crack_geometry": {"length_px": 100},
        "screening_severity": {"level": "high"},
    })

    assert simplified == {
        "index": 2,
        "class_id": 6,
        "class_name": "Structural crack",
        "score": 0.0,
        "detection_confidence": 0.91,
        "screening_severity": {"level": "high"},
    }


def test_docx_preview_reads_text_without_modifying_source(tmp_path: Path) -> None:
    from docx import Document

    import runtime.damage_workflow_gui as gui

    source = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("可读取的损伤报告正文")
    document.save(source)
    original_bytes = source.read_bytes()

    source_path, text = gui.read_document_preview(source)

    assert source_path == str(source)
    assert "可读取的损伤报告正文" in text
    assert source.read_bytes() == original_bytes
